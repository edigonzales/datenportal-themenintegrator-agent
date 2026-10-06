from __future__ import annotations

from pathlib import Path

from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl import load_workbook


def render(root: Path, output: Path, data: dict) -> str:
    env = Environment(loader=FileSystemLoader(root / "templates"), autoescape=select_autoescape(["html"]))
    output.write_text(env.get_template("review.html").render(**data), encoding="utf-8")
    return str(output)


def workbook_cells(path: Path, sheet: str | None = None, offset: int = 0, limit: int = 100) -> dict:
    """Return source coordinates and cached formula results; never execute workbook content."""
    values = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if sheet is None:
            return {"sheets": [{"name": s.title, "rows": s.max_row, "columns": s.max_column} for s in values]}
        ws, fs = values[sheet], formulas[sheet]
        rows = []
        for row, formula_row in zip(
            ws.iter_rows(min_row=offset + 1, max_row=offset + limit),
            fs.iter_rows(min_row=offset + 1, max_row=offset + limit),
        ):
            cells = []
            for cell, formula in zip(row, formula_row):
                if cell.value is None and formula.value is None:
                    continue
                cells.append(
                    {
                        "cell": cell.coordinate,
                        "value": cell.value
                        if isinstance(cell.value, (str, int, float, bool, type(None)))
                        else str(cell.value),
                        "formula": formula.value if formula.data_type == "f" else None,
                        "missing_cached_result": formula.data_type == "f" and cell.value is None,
                    }
                )
            if cells:
                rows.append(cells)
        return {
            "file": path.name,
            "sheet": sheet,
            "rows": rows,
            "next_offset": offset + limit if offset + limit < ws.max_row else None,
        }
    finally:
        values.close()
        formulas.close()
